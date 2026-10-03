"""Builds the sample inputs used by the tests and the README examples.

    python -m tests.samples ./data      # writes every sample into ./data
"""
from __future__ import annotations

import datetime as dt
import io
import sys
import zipfile
from email.message import EmailMessage
from pathlib import Path


def invoice_xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = "ACME CORP - TAX INVOICE"
    ws["A3"] = "Invoice No: INV-20194"
    ws["A4"], ws["B4"] = "Invoice Date", dt.date(2026, 8, 15)
    ws["A5"], ws["B5"] = "Due Date", dt.date(2026, 9, 14)
    ws["A6"] = "Supplier: Acme Corp"
    ws["A7"] = "PO Number: PO-5531"
    ws.append([])
    ws.append(["Description", "Qty", "Unit Price", "Amount"])          # row 9
    ws.append(["Consulting", 10, 1000, 10000])
    ws.append(["Licence", 1, 400, 400])
    ws.append(["Tax", None, None, 2000])
    ws.append(["Total Due", None, None, 12400])                         # row 13
    for r in range(10, 14):
        for c in ("C", "D"):
            ws[f"{c}{r}"].number_format = '"$"#,##0.00'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def statement_csv() -> bytes:
    return (
        "Invoice No,Invoice Date,Due Date,Supplier,PO Number,Total Due\n"
        "INV-30001,15/08/2026,14/09/2026,Globex Ltd,PO-7700,\"$4,250.75\"\n"
    ).encode()


def invoices_csv() -> bytes:
    """A statement listing three invoices, one per row."""
    return (
        "Invoice No,Invoice Date,Due Date,Supplier,Total Due\n"
        "INV-1001,2026-08-01,2026-08-31,Acme Corp,\"$1,250.00\"\n"
        "INV-1002,2026-08-05,2026-09-04,Globex Ltd,\"$980.40\"\n"
        "INV-1003,2026-08-09,2026-09-08,Initech LLC,\"$3,100.00\"\n"
    ).encode()


def invoice_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Table
    buf = io.BytesIO()
    st = getSampleStyleSheet()
    doc = SimpleDocTemplate(buf, pagesize=A4, invariant=1)
    table = Table([["Description", "Qty", "Amount"],
                   ["Hosting", "12", "$1,200.00"],
                   ["Support", "1", "$300.00"],
                   ["Total Due", "", "$1,500.00"]],
                  style=[("GRID", (0, 0), (-1, -1), 0.5, "black")])
    doc.build([Paragraph("Initech Invoice", st["Title"]),
               Paragraph("Invoice No: INI-0042", st["Normal"]),
               Paragraph("Invoice Date: 2026-09-01", st["Normal"]),
               Paragraph("Due Date: 2026-10-01", st["Normal"]),
               Paragraph("Supplier: Initech LLC", st["Normal"]),
               table])
    return buf.getvalue()


def logo_png() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "white").save(buf, "PNG")
    return buf.getvalue()


def invoice_eml() -> bytes:
    msg = EmailMessage()
    msg["From"] = "billing@acme.example"
    msg["To"] = "ap@buyer.example"
    msg["Subject"] = "Invoice INV-20194 from Acme Corp"
    msg["Message-ID"] = "<inv-20194@acme.example>"
    msg.set_content("Hi team,\n\nPlease find invoice INV-20194 attached.\n"
                    "Total due: $12,400.00 by 14 Sep 2026.\n\nThanks,\nAcme billing\n")
    msg.add_attachment(invoice_xlsx(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       filename="INV-20194.xlsx")
    msg.add_attachment(logo_png(), maintype="image", subtype="png", filename="logo.png")
    return bytes(msg)


def invoice_msg() -> bytes:
    """The same invoice email as invoice_eml(), saved from Outlook as .msg."""
    from tests.msg_writer import build_msg
    return build_msg(subject="Invoice INV-20194 from Acme Corp", sender_name="Acme billing",
                     sender_email="billing@acme.example",
                     body="Hi team,\r\n\r\nPlease find invoice INV-20194 attached.\r\n"
                          "Total due: $12,400.00 by 14 Sep 2026.\r\n\r\nThanks,\r\nAcme billing\r\n",
                     attachments=[{"name": "INV-20194.xlsx", "data": invoice_xlsx()},
                                  {"name": "logo.png", "data": logo_png()}])


def bundle_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("statement.csv", statement_csv())
        z.writestr("scans/photo.png", logo_png())
        z.writestr("__MACOSX/._statement.csv", b"junk")
    return buf.getvalue()


SAMPLES = {
    "invoice.xlsx": invoice_xlsx,
    "statement.csv": statement_csv,
    "invoices.csv": invoices_csv,
    "invoice.pdf": invoice_pdf,
    "invoice-email.eml": invoice_eml,
    "invoice-email.msg": invoice_msg,
    "bundle.zip": bundle_zip,
    "logo.png": logo_png,
}


def write_all(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    for name, build in SAMPLES.items():
        (folder / name).write_bytes(build())


if __name__ == "__main__":
    write_all(Path(sys.argv[1] if len(sys.argv) > 1 else "data"))
