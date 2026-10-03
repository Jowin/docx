"""Excel sheets laid out like a document: text view, amounts inside text,
content outside cells (images, charts, text boxes, comments, header/footer,
embedded files)."""
import io
import re
import zipfile

import pytest

from extractor_tools import list_sheets, read_sheet

_TEXTBOX = (
    b'<xdr:twoCellAnchor xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
    b'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><xdr:from><xdr:col>7</xdr:col><xdr:colOff>0</xdr:colOff>'
    b'<xdr:row>9</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:from><xdr:to><xdr:col>10</xdr:col>'
    b'<xdr:colOff>0</xdr:colOff><xdr:row>12</xdr:row><xdr:rowOff>0</xdr:rowOff></xdr:to>'
    b'<xdr:sp macro="" textlink=""><xdr:nvSpPr><xdr:cNvPr id="9" name="TextBox 1"/><xdr:cNvSpPr txBox="1"/>'
    b'</xdr:nvSpPr><xdr:spPr/><xdr:txBody><a:bodyPr/><a:p><a:r><a:t>Bank: HDFC, IFSC HDFC0001234</a:t></a:r></a:p>'
    b'<a:p><a:r><a:t>Late fee EUR 50.00</a:t></a:r></a:p></xdr:txBody></xdr:sp><xdr:clientData/></xdr:twoCellAnchor>'
)


@pytest.fixture
def doc_xlsx() -> bytes:
    import openpyxl
    from openpyxl.chart import BarChart, Reference
    from openpyxl.comments import Comment
    from openpyxl.drawing.image import Image as XImg
    from PIL import Image

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws.merge_cells("B2:F2")
    ws["B2"] = "ACME CORP - TAX INVOICE"
    ws["B4"] = "Invoice No: INV-20194"
    ws["B5"] = "Date: 15 Aug 2026"
    ws.merge_cells("B7:F9")
    ws["B7"] = "Payment is due within 30 days. Please remit to account ending 4411."
    ws["E12"] = "Total Due: $12,400.00"
    ws["B14"], ws["C14"] = "Tax", 0.18
    ws["C14"].number_format = "0%"
    ws["B15"], ws["C15"] = "Freight", 250
    ws["C15"].number_format = '"$"#,##0.00'
    ws["B16"] = "Old total: $99,999.00"
    ws.row_dimensions[16].hidden = True
    ws["H20"] = "x"
    ws["H20"].comment = Comment("PO-5531 approved", "AP clerk")
    ws.oddFooter.center.text = "Acme Corp Confidential"
    png = io.BytesIO()
    Image.new("RGB", (200, 80), "white").save(png, "PNG")
    png.seek(0)
    ws.add_image(XImg(png), "H2")
    data = wb.create_sheet("Data")
    for i in range(1, 4):
        data.append([f"M{i}", i * 10])
    chart = BarChart()
    chart.add_data(Reference(data, min_col=2, min_row=1, max_row=3))
    data.add_chart(chart, "D2")
    buf = io.BytesIO()
    wb.save(buf)

    # Add what openpyxl cannot write: a text box and an embedded PDF.
    src, out = zipfile.ZipFile(io.BytesIO(buf.getvalue())), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.infolist():
            body = src.read(item.filename)
            if item.filename == "xl/drawings/drawing1.xml":
                close = b"</xdr:wsDr>" if b"</xdr:wsDr>" in body else b"</wsDr>"
                assert close in body
                body = body.replace(close, _TEXTBOX + close)
            if item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                body = body.replace(b"</Relationships>",
                                    b'<Relationship Id="rIdPkg" Type="http://schemas.openxmlformats.org/'
                                    b'officeDocument/2006/relationships/package" '
                                    b'Target="../embeddings/remittance.pdf"/></Relationships>')
            z.writestr(item, body)
        z.writestr("xl/embeddings/remittance.pdf", b"%PDF-1.4\n% remittance advice\n%%EOF\n")
    return out.getvalue()


def test_inventory_sees_a_document_sheet(doc_xlsx):
    r = list_sheets(doc_xlsx)["result"]
    inv = r["sheets"][0]
    assert inv["currency_text_cells"] == 2          # "$12,400.00" and the hidden "$99,999.00"
    assert inv["currency_density"] > 0             # was 0.0 before in-string amounts
    assert inv["label_value_cells"] == 4
    assert inv["long_text_cells"] == 1
    assert inv["outside_cells"] == {"inspected": True, "images": 1, "charts": 0, "text_boxes": 1,
                                    "comments": 1, "embedded": 1, "header_footer": True}
    assert r["sheets"][1]["outside_cells"]["charts"] == 1
    emb = r["embedded_objects"][0]
    assert emb["path"] == "xl/embeddings/remittance.pdf" and emb["detected"] == "pdf"


def test_amounts_inside_string_cells(doc_xlsx):
    rows = read_sheet(doc_xlsx, sheet="Invoice")["result"]["rows"]
    cell = {c["ref"]: c for row in rows for c in row["cells"]}["E12"]
    a = cell["amounts"][0]
    assert (a["value"], a["currency"], a["locator"]) == ("12400.00", "$", "E12:C12-21")


def test_text_view_reads_like_a_document(doc_xlsx):
    r = read_sheet(doc_xlsx, sheet="Invoice", view="text")["result"]
    assert "rows" not in r
    lines = {ln["row"]: ln for ln in r["text_lines"]}
    assert lines[2]["text"] == "B2: ACME CORP - TAX INVOICE"
    assert lines[14]["text"] == "B14: Tax | C14: 18%"
    assert lines[15]["amounts"] == [{"text": "$250", "value": "250", "currency": "$", "locator": "C15"}]
    assert lines[12]["amounts"][0]["locator"] == "E12:C12-21"
    assert 16 not in lines                           # hidden row stays out of the text view


def test_outside_cell_content_is_returned(doc_xlsx):
    out = read_sheet(doc_xlsx, sheet="Invoice", view="text")["result"]["outside_cells"]
    assert out["comments"] == [{"ref": "H20", "text": "PO-5531 approved", "author": "AP clerk"}]
    assert out["text_boxes"] == [{"anchor": "H10", "text": "Bank: HDFC, IFSC HDFC0001234\nLate fee EUR 50.00"}]
    assert out["images"] == [{"anchor": "H2"}]
    assert out["header_footer"] == {"oddFooter": "Acme Corp Confidential"}
    assert out["embedded"] == [{"path": "xl/embeddings/remittance.pdf"}]


def test_both_view(doc_xlsx):
    r = read_sheet(doc_xlsx, sheet="Invoice", view="both")["result"]
    assert r["rows"] and r["text_lines"]


def test_xls_says_outside_cells_not_inspected(invoice_xls):
    r = read_sheet(invoice_xls, sheet="Transactions", view="text")["result"]
    assert r["outside_cells"] == {"inspected": False}
    assert r["text_lines"][1]["text"].startswith("A2: INV-20194 | B2: $12400.5")
