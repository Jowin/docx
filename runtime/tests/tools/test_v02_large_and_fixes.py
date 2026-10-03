"""v0.2: paging, tails, Excel precision, locale tags, hidden rows/columns."""
import io
import re
import zipfile

import pytest

from extractor_tools import ToolError, list_sheets, read_csv, read_sheet
from extractor_tools.spreadsheet_reader import currency_of_format


def _xlsx(build) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    build(wb.active)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _cells(rows):
    return {c["ref"]: c for row in rows for c in row["cells"]}


# ---------------------------------------------------------------- CSV paging

@pytest.fixture
def long_csv() -> bytes:
    lines = ["Ref,Amount"] + [f"R{i},{i}.25" for i in range(1, 1001)] + ["Total,\"500,750.00\""]
    return "\n".join(lines).encode()


def test_csv_first_page_and_tail(long_csv):
    r = read_csv(long_csv, row_limit=100)["result"]
    assert r["data_row_count"] == 1001 and len(r["rows"]) == 100
    assert r["page"] == {"row_offset": 0, "row_limit": 100, "returned": 100,
                         "has_more": True, "next_offset": 100}
    assert r["tail"][-1]["cells"][1]["value"] == "500750.00"       # the total, on page 1
    assert r["columns"][1]["type_counts"] == {"decimal": 1001}      # profile covers every row


def test_csv_last_page_does_not_repeat_tail(long_csv):
    r = read_csv(long_csv, row_offset=1000, row_limit=100)["result"]
    assert [row["index"] for row in r["rows"]] == [1000]
    assert r["page"]["has_more"] is False and r["page"]["next_offset"] is None
    assert 1000 not in [row["index"] for row in r["tail"]]


def test_csv_profile_only(long_csv):
    r = read_csv(long_csv, row_limit=0, tail_rows=0)["result"]
    assert r["rows"] == [] and r["tail"] == [] and r["columns"][0]["name"] == "Ref"


def test_csv_unicode_minus():
    r = read_csv("Amt\n−450.50\n".encode())["result"]
    assert r["rows"][0]["cells"][0]["value"] == "-450.50"


def test_csv_bad_page_params(long_csv):
    with pytest.raises(ToolError) as e:
        read_csv(long_csv, row_limit=10_000)
    assert e.value.code == "param_invalid"


# ---------------------------------------------------------------- Excel fixes

def test_float_noise_reported_at_excel_precision():
    def build(ws):
        ws["A1"] = 12400.499999999998
        ws["A2"] = 0.1 + 0.2
        ws["A3"] = 1e-7
    c = _cells(read_sheet(_xlsx(build), sheet="Sheet")["result"]["rows"])
    assert c["A1"]["value"] == "12400.5"
    assert c["A2"]["value"] == "0.3"
    assert c["A3"]["value"] == "0.0000001"


@pytest.mark.parametrize("fmt,expected", [
    ("[$-409]mmmm d, yyyy", None),            # locale tag only
    ("[$-409]#,##0.00", None),
    ("[$€-407]#,##0.00", "€"),
    ('[$USD] #,##0.00', "USD"),
    ('"$"#,##0.00', "$"),
    ("#,##0.00 [$£-809]", "£"),
    ("0.00%", None),
    ("General", None),
])
def test_currency_of_format(fmt, expected):
    assert currency_of_format(fmt) == expected


def test_locale_tagged_number_is_not_currency():
    def build(ws):
        ws["A1"] = 42
        ws["A1"].number_format = "[$-409]#,##0"
        ws["A2"] = 42
        ws["A2"].number_format = "[$€-407]#,##0.00"
    c = _cells(read_sheet(_xlsx(build), sheet="Sheet")["result"]["rows"])
    assert "currency_format" not in c["A1"]
    assert c["A2"]["currency_symbol"] == "€"


def test_hidden_rows_and_columns_are_marked():
    def build(ws):
        for r in range(1, 5):
            ws.cell(row=r, column=1, value=r)
            ws.cell(row=r, column=3, value=r * 10)
        ws.row_dimensions[3].hidden = True
        ws.column_dimensions["C"].hidden = True
    data = _xlsx(build)
    r = read_sheet(data, sheet="Sheet")["result"]
    c = _cells(r["rows"])
    assert c["A3"]["hidden"] and c["C1"]["hidden"] and "hidden" not in c["A1"]
    assert r["hidden_rows"]["items"] == [3] and r["hidden_columns"] == [3]
    inv = list_sheets(data)["result"]["sheets"][0]
    assert inv["hidden_rows"] == 1 and inv["hidden_columns"] == 1


def test_merged_top_left_carries_range(invoice_xlsx):
    c = _cells(read_sheet(invoice_xlsx, sheet="Summary")["result"]["rows"])
    assert c["A3"]["merged_range"] == "A3:C3"


# ---------------------------------------------------------------- Excel paging

@pytest.fixture
def long_xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Lines")
    ws.append(["Item", "Amount"])
    for i in range(1, 1201):
        ws.append([f"Line {i}", i * 1.25])
    ws.append(["Total", 900750.0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_pages_and_tail(long_xlsx):
    r = read_sheet(long_xlsx, sheet="Lines", row_limit=500)["result"]
    assert r["row_count"] == 1202 and r["page"]["next_start_row"] == 501
    assert len(r["rows"]) == 500
    assert r["tail"][-1]["cells"][0]["value"] == "Total"
    last = read_sheet(long_xlsx, sheet="Lines", start_row=1001, row_limit=500)["result"]
    assert last["page"]["has_more"] is False and len(last["rows"]) == 202
    assert last["tail"] == []                                      # already on the page


def test_wrong_dimension_tag_does_not_truncate(long_xlsx):
    buf_in, buf_out = io.BytesIO(long_xlsx), io.BytesIO()
    with zipfile.ZipFile(buf_in) as zin, zipfile.ZipFile(buf_out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            body = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/"):
                body, n = re.subn(rb'<dimension ref="[^"]*"/>', b'<dimension ref="A1"/>', body)
                if not n:                               # write-only mode omits it; add a wrong one
                    body, n = re.subn(rb"<sheetViews", b'<dimension ref="A1"/><sheetViews', body, 1)
                if not n:
                    body, n = re.subn(rb"<sheetData", b'<dimension ref="A1"/><sheetData', body, 1)
                assert n == 1
            zout.writestr(item, body)
    r = read_sheet(buf_out.getvalue(), sheet="Lines", row_limit=10, tail_rows=1)["result"]
    assert r["row_count"] == 1202 and r["tail"][0]["row"] == 1202


def test_list_sheets_samples_large_sheets(long_xlsx):
    s = list_sheets(long_xlsx, sample_rows=100)["result"]["sheets"][0]
    assert s["sampled_rows"] == 100 and s["sample_complete"] is False
    assert s["row_count"] == 1202


def test_range_read_has_no_paging(invoice_xlsx):
    r = read_sheet(invoice_xlsx, sheet="Summary", cell_range="A5:B7")["result"]
    assert r["page"] is None and r["tail"] == []


# ---------------------------------------------------------------- XLS

def test_xls_hidden_row_and_precision():
    import xlwt
    wb = xlwt.Workbook()
    s = wb.add_sheet("T")
    s.write(0, 0, 0.1 + 0.2)
    s.write(1, 0, "hidden")
    s.row(1).hidden = True
    buf = io.BytesIO()
    wb.save(buf)
    r = read_sheet(buf.getvalue(), sheet="T")["result"]
    c = _cells(r["rows"])
    assert c["A1"]["value"] == "0.3" and c["A2"]["hidden"]


def test_csv_default_page_is_50_with_full_profile(long_csv):
    r = read_csv(long_csv)["result"]
    assert len(r["rows"]) == 50 and r["page"]["next_offset"] == 50
    assert len(r["tail"]) == 5
    assert r["columns"][1]["non_empty"] == 1001          # profile still covers every row


def test_excel_default_page_is_50(long_xlsx):
    r = read_sheet(long_xlsx, sheet="Lines")["result"]
    assert len(r["rows"]) == 50 and r["page"]["next_start_row"] == 51
    assert r["tail"][-1]["cells"][0]["value"] == "Total"
