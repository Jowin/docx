"""Writer tools: an extraction result -> CSV, Excel, Word, PDF. Pure and deterministic."""
from __future__ import annotations

import csv
import io
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from extractor_tools import ToolError, write, write_csv, write_docx, write_pdf, write_xlsx
from extractor_tools.api import app

RESULT = {
    "data": [
        {"invoice_number": "INV-1", "vendor": "=HYPERLINK(\"x\")", "total_amount": 12.5,
         "line_items": [{"description": "Widget", "amount": 10}, {"description": "Tax", "amount": 2.5}]},
        {"invoice_number": "INV-2", "vendor": "Acme ₹ Ltd", "total_amount": -3.0, "line_items": []},
    ],
    "flagged": True, "confidence": 0.7, "flags": ["low_confidence"],
    "records": [
        {"flagged": False, "confidence": 0.9, "flags": [],
         "fields": {"invoice_number": {"value": "INV-1", "confidence": 0.9, "source": "attachment:a.csv#A2"}}},
        {"flagged": True, "confidence": 0.7, "flags": ["low_confidence"], "fields": {}},
    ],
    "metadata": {"audit_id": "run_abc", "config": {"client": "acme", "usecase": "ap", "version": "1.0.0"},
                 "input": {"name": "m.eml"}},
}


def _rows(data: bytes) -> list[list[str]]:
    return list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))


def test_csv_explodes_the_single_array_and_neutralises_formulas():
    rows = _rows(write_csv(RESULT))
    assert rows[0] == ["_record", "invoice_number", "vendor", "total_amount", "line_items.description",
                       "line_items.amount", "_confidence", "_flagged", "_flags"]
    assert rows[1][:6] == ["1", "INV-1", "'=HYPERLINK(\"x\")", "12.5", "Widget", "10"]
    assert rows[2][4:6] == ["Tax", "2.5"]
    assert rows[3] == ["2", "INV-2", "Acme ₹ Ltd", "-3.0", "", "", "0.7", "true", "low_confidence"]
    assert write_csv(RESULT).startswith(b"\xef\xbb\xbf")


def test_csv_without_explosion_and_from_plain_records():
    rows = _rows(write_csv(RESULT, explode=""))
    assert rows[0][:5] == ["_record", "invoice_number", "vendor", "total_amount", "line_items"]
    assert len(rows) == 3 and rows[1][4].startswith('[{"amount": 10')
    plain = _rows(write_csv(RESULT["data"], bom=False, delimiter=";"))
    assert "_flags" not in plain[0] and len(plain) == 4


def test_every_writer_is_deterministic():
    for fmt in ("csv", "xlsx", "docx", "pdf"):
        assert write(RESULT, fmt)[0] == write(RESULT, fmt)[0], fmt


def test_xlsx_sheets_and_typed_cells():
    wb = load_workbook(io.BytesIO(write_xlsx(RESULT)))
    assert wb.sheetnames == ["Records", "line_items", "Sources", "Run"]
    rec = wb["Records"]
    assert [c.value for c in rec[1]] == ["_record", "invoice_number", "vendor", "total_amount",
                                         "Confidence", "Flagged", "Flags"]
    assert rec["D2"].value == 12.5 and rec["C2"].value.startswith("'=")
    assert [c.value for c in wb["line_items"][3]] == [1, "Tax", 2.5]
    assert wb["Sources"]["E2"].value == "attachment:a.csv#A2"


def test_docx_and_pdf_reports():
    doc = Document(io.BytesIO(write_docx(RESULT)))
    texts = [p.text for p in doc.paragraphs]
    assert "Record 1" in texts and "Flags: low_confidence" in texts
    cells = {c.text for t in doc.tables for r in t.rows for c in r.cells}
    assert {"INV-1", "attachment:a.csv#A2", "Widget"} <= cells
    pdf = write_pdf(RESULT, title="Acme invoices")
    assert pdf.startswith(b"%PDF") and len(pdf) > 1000
    from pypdf import PdfReader
    text = PdfReader(io.BytesIO(pdf)).pages[0].extract_text()
    assert "Acme invoices" in text and "INV-1" in text and "Widget" in text


def test_an_empty_result_still_writes():
    empty = {"data": [], "flagged": True, "flags": ["out_of_scope"], "records": [], "metadata": {}}
    assert _rows(write_csv(empty))[0] == ["_record", "_confidence", "_flagged", "_flags"]
    for fmt in ("xlsx", "docx", "pdf"):
        assert write(empty, fmt)[0]


def test_bad_input_and_format():
    with pytest.raises(ToolError):
        write({"nope": 1}, "csv")
    with pytest.raises(ToolError):
        write(RESULT, "rtf")
    assert write(RESULT, "Excel")[2] == "xlsx" and write(RESULT, "word")[2] == "docx"


def test_writer_endpoints():
    c = TestClient(app)
    r = c.post("/tools/csv-writer", json={"result": RESULT, "params": {"explode": ""}, "filename": "inv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert 'filename="inv.csv"' in r.headers["content-disposition"]
    for path, magic in (("excel-writer", b"PK"), ("docx-writer", b"PK"), ("pdf-writer", b"%PDF")):
        r = c.post(f"/tools/{path}", json={"result": RESULT})
        assert r.status_code == 200 and r.content.startswith(magic), path
    assert zipfile.is_zipfile(io.BytesIO(c.post("/tools/excel-writer", json={"result": RESULT}).content))
    bad = c.post("/tools/pdf-writer", json={"result": RESULT, "params": {"explode": "x"}})
    assert bad.status_code == 422 and bad.json()["error"] == "param_invalid"
