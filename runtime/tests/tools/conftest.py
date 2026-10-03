"""Fixtures are built in code so the tests carry no binary files."""
from __future__ import annotations

import datetime as dt
import io

import pytest


@pytest.fixture
def invoice_csv() -> bytes:
    return (
        "Invoice No,Invoice Date,Due Date,Description,Qty,Unit Price,Total Due\n"
        "INV-20194,15/08/2026,14/09/2026,Consulting,10,\"$1,240.00\",\"$12,400.00\"\n"
        "INV-20195,03/09/2026,03/10/2026,Licence,1,\"$450.50\",\"($450.50)\"\n"
        "\n"
        "00123,,,Total,,,\"$11,949.50\"\n"
    ).encode("utf-8")


@pytest.fixture
def invoice_xlsx() -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Notes"
    ws["A1"] = "Remittance notes only"
    s = wb.create_sheet("Summary")
    s["A1"], s["B1"] = "Invoice", "INV-20194"
    s.merge_cells("A3:C3")
    s["A3"] = "Line items"
    s["A5"], s["B5"], s["C5"] = "Item", "Amount", "Tax rate"
    s["A6"], s["B6"], s["C6"] = "Consulting", 12400, 0.2
    s["B6"].number_format = '"$"#,##0.00'
    s["C6"].number_format = "0%"
    s["A7"], s["B7"] = "Discount", 49.5
    s["B7"].number_format = '"$"#,##0.00'
    s["A14"], s["B14"] = "Total Due", "=B6-B7"   # no cached value: never opened in Excel
    s["A15"], s["B15"] = "Due", dt.date(2026, 9, 15)
    s["A16"], s["B16"] = "Paid", False
    h = wb.create_sheet("Hidden")
    h.sheet_state = "hidden"
    h["A1"] = 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def invoice_xls() -> bytes:
    import xlwt
    wb = xlwt.Workbook()
    s = wb.add_sheet("Transactions")
    money = xlwt.easyxf(num_format_str='"$"#,##0.00')
    datef = xlwt.easyxf(num_format_str="YYYY-MM-DD")
    s.write(0, 0, "Invoice")
    s.write(0, 1, "Amount")
    s.write(0, 2, "Due")
    s.write(1, 0, "INV-20194")
    s.write(1, 1, 12400.5, money)
    s.write(1, 2, dt.date(2026, 9, 15), datef)
    s.write_merge(3, 3, 0, 2, "Footer")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def invoice_pdf() -> bytes:
    from reportlab.lib.pagesizes import A4
    from PIL import Image as PILImage
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, PageBreak, Image
    from reportlab.lib.styles import getSampleStyleSheet
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, invariant=1)
    st = getSampleStyleSheet()
    table = Table([["Item", "Amount"], ["Consulting", "$12,400.00"], ["Total Due", "$12,400.00"]],
                  style=[("GRID", (0, 0), (-1, -1), 0.5, "black")])
    png = io.BytesIO()
    PILImage.new("L", (400, 560), 200).save(png, "PNG")
    png.seek(0)
    doc.build([Paragraph("Invoice INV-20194", st["Title"]),
               Paragraph("Payment due 2026-09-15. Remit to Acme Corp.", st["Normal"]),
               table, PageBreak(),
               Image(png, width=400, height=560)])   # page 2 simulates a scan: image, no text
    return buf.getvalue()


def encrypt_pdf(data: bytes, user_pw: str, owner_pw: str) -> bytes:
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
    w.encrypt(user_password=user_pw, owner_password=owner_pw, algorithm="AES-128")
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()
